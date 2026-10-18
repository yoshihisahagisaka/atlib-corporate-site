"""Fixed comparison schedules; intentionally no formulas, input changes do not refresh."""
import json
from pathlib import Path
from openpyxl.styles import Font,PatternFill,Alignment
def add_sheets(wb):
    d=json.loads(Path('secure-services-business-analysis/docs/secure-services-20261018/c40-c50/c-results.json').read_text(encoding='utf8'))
    schedules=[('価格と高利用感度',['価格定義','高利用率','運用費含む','DNS最低額','構成','配分率','初回黒字社','継続黒字社','両サービス黒字社','平均売上','平均配分','100社G差引','100社S差引','10社合計','30社','50社','100社','300社','1000社'],[[r['price'],r['h'],int(r['support']),r['dns'],d['labels'][r['mi']],r['share'],r['first'],r['sustained'],r['both'],r['avgRevenue'],r['avgShare'],r['gProfit100'],r['sProfit100'],*r['profits']] for r in d['sensitivities']]),('獲得速度比較',['構成','配分率','新規社/月','単月初回黒字月','累積回収月','最大累積赤字','36月累積差引'],[[d['labels'][r['mi']],r['share'],r['pace'],r['first'] or '36月内なし',r['sustainedRecovery'] or '36月内なし',r['maxFunding'],r['cumulative36']] for r in d['growth']])]
    for name,headers,rows in schedules:
        w=wb.create_sheet(name);w.sheet_view.showGridLines=False;w.append([name+'（作成時固定表）']);w.append(['入力変更に連動しない。round=元価格 / service20=各サービス基本20円減 / product20=パックを含む各商品20円減。floor=MAX / add=最低額加算。運用費含む0は固定運用とサポートを除外。']);w.append([]);w.append(headers)
        for r in rows:w.append(r)
        w.merge_cells(start_row=1,start_column=1,end_row=1,end_column=min(9,len(headers)));w.merge_cells(start_row=2,start_column=1,end_row=2,end_column=min(9,len(headers)))
        w.row_dimensions[1].height=32;w.row_dimensions[2].height=65;w.row_dimensions[4].height=50
        for row in w:
            for cell in row:cell.font=Font(name='Yu Gothic',size=10,color='203047');cell.alignment=Alignment(vertical='center',wrap_text=True);cell.number_format='#,##0;(#,##0);"-"'
        for cell in w[4]:cell.fill=PatternFill('solid',fgColor='1F3556');cell.font=Font(name='Yu Gothic',size=10,bold=True,color='FFFFFF')
        for col in range(1,len(headers)+1):w.column_dimensions[w.cell(4,col).column_letter].width=22 if col in [1,5] else 18
        w.freeze_panes='B5';w.auto_filter.ref=f'A4:{w.cell(4,len(headers)).column_letter}{w.max_row}'
        for row in range(5,w.max_row+1):
            w.row_dimensions[row].height=32
            for col in ([2,6] if name=='価格と高利用感度' else [2]):w.cell(row,col).number_format='0%'
