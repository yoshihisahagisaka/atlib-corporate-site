"""Offline table previews of saved cached values, no native Excel rendering."""
from pathlib import Path
from openpyxl import load_workbook
from PIL import Image,ImageDraw,ImageFont
import math,json
out=Path('outputs/secure-services-c40-c50-20261018');w=load_workbook(out/'Secure_Services_C40_C50_Model.xlsx',data_only=True)
font=ImageFont.truetype('C:/Windows/Fonts/YuGothM.ttc',17);small=ImageFont.truetype('C:/Windows/Fonts/YuGothM.ttc',15)
images=[]
for idx,s in enumerate(w):
    # First output/input block on each tab. Other columns checked structurally.
    count=min(9,s.max_column);cw=190;im=Image.new('RGB',(count*cw+40,820),'white');dr=ImageDraw.Draw(im)
    dr.text((20,8),s.title,fill='#1f3556',font=font)
    for row in range(1,11):
        y=38+(row-1)*75
        for col in range(1,count+1):
            x=20+(col-1)*cw;v=s.cell(row,col).value
            if v is None:continue
            if isinstance(v,(int,float)):text=f'{v:,.0f}' if abs(v)>1 else str(round(v,3))
            else:text=str(v)
            lines=[];line=''
            for char in text:
                if dr.textlength(line+char,font=small)>cw-12:lines.append(line);line=''
                line+=char
            lines.append(line)
            if len(lines)>3:lines=lines[:2]+[lines[2][:-1]+'…']
            dr.rectangle((x,y,x+cw,y+75),outline='#d8e0e8',fill='#eef4f8' if row==4 else 'white')
            dr.multiline_text((x+5,y+5),'\n'.join(lines),font=small,fill='#203047',spacing=3)
    path=out/f'preview-{idx+1:02}.png';im.save(path);images.append(im)
contact=Image.new('RGB',(1800,math.ceil(len(images)/3)*300),'#d8e0e8')
for idx,im in enumerate(images):im.thumbnail((595,295));contact.paste(im,((idx%3)*600,(idx//3)*300))
contact.save(out/'preview-contact.png')
# Structural QA and cache readback.
f=load_workbook(out/'Secure_Services_C40_C50_Model.xlsx',data_only=False)
errors=[]
for s in w:
    for row in s:
        for c in row:
            if c.data_type=='e':errors.append((s.title,c.coordinate,c.value))
assert not errors
assert w['前提と出典']['B17'].value==1000
assert w['前提と出典']['C42'].value==3500
assert w['経営比較']['C8'].value==22 and w['経営比較']['D8'].value==55
assert w['経営比較']['C9'].value==25 and w['経営比較']['D9'].value==62
assert abs(w['経営比較']['E8'].value-197167.98665653646)<1e-7
assert abs(w['経営比較']['E9'].value-164724.58529217273)<1e-7
assert f['経営比較']['E8'].data_type=='f'
(out/'c-workbook-verification.json').write_text(json.dumps({'sheets':len(w.sheetnames),'selectedC40BE':22,'selectedC50BE':25,'bothC40':55,'bothC50':62,'errorCells':errors,'preview':'offline cached-value previews, not native Excel render','nativeExcelRecalculation':'unverified'},ensure_ascii=False,indent=2),encoding='utf8')
print('Verified C workbook values and formulas; generated nine offline previews.')
