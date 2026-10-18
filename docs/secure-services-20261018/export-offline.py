"""Offline XLSX authoring and cache verification. No network or analytics imports."""
import json, math, zipfile, os
from pathlib import Path
from xml.etree import ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import range_boundaries, get_column_letter

out=Path('outputs/secure-services-integrated-20261018')
data=json.loads((out/'offline-workbook-data.json').read_text(encoding='utf8'))
cache=json.loads((out/'offline-value-cache.json').read_text(encoding='utf8'))
wb=Workbook();wb.remove(wb.active)
headers={'経営比較':[6,12,21,35],'前提と出典':[35,40,45,51,66,72,79],'同一企業モデル':[4],'契約別単価':[4],'利益曲線':[4],'累積回収':[4],'固定条件比較':[4]}
formula_count=0
for info in data:
    ws=wb.create_sheet(info['name']);ws.sheet_view.showGridLines=False
    for a,v in info['cells'].items():
        cell=ws[a];value=v['value']
        if isinstance(value,str) and not v['formula']:
            value=value.replace('G0現行基準','G0基本5000比較').replace('G0現行','G0基本5000比較')
            if ws.title=='経営比較' and a=='A2':value='選択条件は4行に表示。共通率は前提と出典E4〜E6。価格・配分は正式決定ではない。'
            if ws.title=='経営比較':value=value.replace('40%配分','率2配分').replace('40%後差引','率2後差引')
            if ws.title in ['利益曲線','累積回収']:
                for rate,idx in [('30','1'),('40','2'),('50','3')]:value=value.replace(rate+'%後','・率'+idx+'後').replace('累積'+rate,'累積・率'+idx).replace('資金谷'+rate,'資金谷・率'+idx).replace('合計'+rate,'合計・率'+idx).replace('G'+rate,'G・率'+idx).replace('S'+rate,'S・率'+idx).replace('初黒'+rate,'初黒・率'+idx).replace('最終赤'+rate,'最終赤・率'+idx).replace('両黒'+rate,'両黒・率'+idx).replace('単黒'+rate+'月','単黒・率'+idx+'月').replace('累回'+rate+'月','累回・率'+idx+'月')
        cell.value=value;cell.font=Font(name='Arial',size=10,color='008000' if v['formula'] and "'" in str(v['value']) and ws.title!='経営比較' else '203047');cell.alignment=Alignment(vertical='center',horizontal='right' if v['formula'] or isinstance(v['value'],(int,float)) else 'left')
        if v['formula']:
            formula_count+=1
            if a not in cache.get(ws.title,{}): raise ValueError(f'Missing cache: {ws.title}!{a}')
    for a,fmt in info['formats']:
        for row in (ws[a] if ':' in a else [[ws[a]]]):
            for c in row:c.number_format=fmt
    for col in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=19
    for row in range(1,ws.max_row+1):ws.row_dimensions[row].height=24
    for row in headers.get(ws.title,[]):
        ws.row_dimensions[row].height=42
        for cell in ws[row]:cell.fill=PatternFill('solid',fgColor='1F3556');cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True,vertical='center')
    ws['A1'].font=Font(name='Arial',size=15,bold=True,color='1F3556');ws.row_dimensions[1].height=32
    ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=min(ws.max_column,9))
    ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=min(ws.max_column,9));ws['A2'].alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[2].height=48
    if info['freeze']:ws.freeze_panes='B'+str(info['freeze']+1)
    if ws.title=='経営比較':
        ws.column_dimensions['A'].width=30
        for c in ['B','C','D']:ws.column_dimensions[c].width=25
        for row in [31,32,33]:ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=9);ws.cell(row,1).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[row].height=38
        for row in [36,37,38]:ws.cell(row,1).number_format='0%';ws.cell(row,2).number_format='#,##0;(#,##0);"-"';ws.cell(row,3).number_format='#,##0;(#,##0);"-"'
    if ws.title=='前提と出典':
        ws.column_dimensions['A'].width=44;ws.column_dimensions['B'].width=25
        for a in ['B4:B38','B41:G43','B46:I48','B52:E58','E4:E6','E10:E12']:
            for row in ws[a]:
                for cell in row:
                    if cell.data_type!='f' and isinstance(cell.value,(int,float)):cell.font=Font(name='Arial',size=10,color='0000FF');cell.fill=PatternFill('solid',fgColor='FFF2CC')
        for row in [60,85,86,87,88,89,90]:ws.merge_cells(start_row=row,start_column=1,end_row=row,end_column=9);ws.cell(row,1).alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[row].height=38
        for row,maximum in [(14,2),(15,2),(16,2),(18,3),(19,1),(20,1),(33,1),(34,3)]:
            dv=DataValidation(type='whole',operator='between',formula1=0,formula2=maximum);dv.errorTitle='入力範囲';dv.error='指定された番号を入力してください';dv.showErrorMessage=True;ws.add_data_validation(dv);dv.add('B'+str(row))
    if ws.title in ['経営比較','利益曲線','累積回収']:
        bounds={'経営比較':'E7:G9','利益曲線':'R5:Z1004','累積回収':'S5:AG40'}[ws.title]
        ws.conditional_formatting.add(bounds,CellIsRule(operator='lessThan',formula=['0'],fill=PatternFill('solid',fgColor='FDE9E7'),font=Font(color='B71C1C')))
wb.calculation=CalcProperties(calcId=191029,fullCalcOnLoad=True,forceFullCalc=True,calcMode='auto')
target=out/'Secure_Services_Integrated_Model.xlsx';wb.save(target)
ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'};n=ns['m']
tmp=out/'cache-patched.xlsx'
with zipfile.ZipFile(target) as src,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as dst:
    for item in src.infolist():
        b=src.read(item.filename)
        if item.filename.startswith('xl/worksheets/sheet') and item.filename.endswith('.xml'):
            idx=int(item.filename.split('sheet')[-1].split('.')[0])-1;name=wb.sheetnames[idx];tree=ET.fromstring(b)
            for c in tree.findall('.//m:c',ns):
                if c.find('m:f',ns) is None:continue
                value=cache[name][c.attrib['r']];v=c.find('m:v',ns)
                if v is None:v=ET.SubElement(c,'{'+n+'}v')
                if isinstance(value,str):c.set('t','str');v.text=value
                else:c.attrib.pop('t',None);v.text=str(value)
            b=ET.tostring(tree,encoding='utf-8',xml_declaration=True)
        dst.writestr(item,b)
os.replace(tmp,target)
values=load_workbook(target,data_only=True);formulas=load_workbook(target,data_only=False)
assert values['経営比較']['C8'].value==25
assert values['経営比較']['D8'].value==53
assert abs(values['経営比較']['E8'].value-166147.98665653652)<1e-6
assert formulas['経営比較']['E8'].data_type=='f'
assert len(wb.sheetnames)==7
note={'export':'offline openpyxl plus OOXML numeric caches','network':'none in generation scripts','formulaCells':formula_count,'cachedValues':'independent JS model; earlier authored formulas verified with Artifact Tool before blocked export','selectedBE40':25,'bothServicesBE40':53,'nativeExcelRecalculation':'not verified','reason':'Automatic approval review rejected initial export after detecting an external Visual Studio telemetry request with unknown payload. Initial process ended; no bypass of that exporter was attempted.'}
(out/'offline-export-verification.json').write_text(json.dumps(note,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps({'saved':str(target),'formulaCells':formula_count,'sheets':len(wb.sheetnames)},ensure_ascii=False))
